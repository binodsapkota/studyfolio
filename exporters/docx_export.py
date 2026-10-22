from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document

from agent.schema import StudyDocument


def docx_bytes(document: StudyDocument) -> bytes:
    word = Document()
    if document.show_title:
        word.add_heading(document.title, level=0)
    if document.author_note.strip():
        note = word.add_paragraph(document.author_note)
        note.runs[0].italic = True
    for section in document.sections:
        if section.heading.strip():
            word.add_heading(section.heading, level=1)
        for paragraph in section.paragraphs:
            word.add_paragraph(paragraph)
        for bullet in section.bullets:
            word.add_paragraph(bullet, style="List Bullet")
    if document.references:
        word.add_heading("References", level=1)
        for reference in document.references:
            word.add_paragraph(reference)
    buffer = BytesIO()
    word.save(buffer)
    return buffer.getvalue()


def write_docx(document: StudyDocument, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(docx_bytes(document))
    return path
