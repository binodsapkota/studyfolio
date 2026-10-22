from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from agent.schema import WorkbookModel


def xlsx_bytes(workbook: WorkbookModel) -> bytes:
    book = Workbook()
    used: set[str] = set()
    sheets = workbook.sheets or []
    if not sheets:
        book.active.title = "Sheet1"
    for index, sheet in enumerate(sheets):
        worksheet = book.active if index == 0 else book.create_sheet()
        worksheet.title = _unique_title(sheet.name, used)
        for column, header in enumerate(sheet.headers, start=1):
            worksheet.cell(1, column, _coerce(header))
        for row_index, row in enumerate(sheet.rows, start=2):
            for column, value in enumerate(row, start=1):
                worksheet.cell(row_index, column, _coerce(value))
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def write_xlsx(workbook: WorkbookModel, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(xlsx_bytes(workbook))
    return path


def _coerce(value: str):
    text = str(value).strip()
    if not text or text.startswith("="):
        return text
    try:
        if any(mark in text for mark in ".eE"):
            return float(text)
        return int(text)
    except ValueError:
        return text


def _unique_title(name: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", name).strip() or "Sheet"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate in used:
        tail = f" {suffix}"
        candidate = base[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate)
    return candidate
